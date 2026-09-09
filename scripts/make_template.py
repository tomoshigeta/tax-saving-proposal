"""物件データ入力テンプレート(.xlsx)を生成する。

列定義は src/storage/excelColumns.json を唯一の正とし、
アプリ側の取り込み処理と同じ定義から作る。

    python3 scripts/make_template.py
"""

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
SPEC = json.loads((ROOT / "src/storage/excelColumns.json").read_text(encoding="utf-8"))
OUT = ROOT / "public/rent-assessment-template.xlsx"

FONT = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="E8EEF7")
REQUIRED_FILL = PatternFill("solid", fgColor="FFF3CD")   # 必須列の見出し
EXAMPLE_FILL = PatternFill("solid", fgColor="F4F3F0")
THIN = Side(style="thin", color="C4C2BA")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

YEN_FMT = "#,##0"
PCT_FMT = "0.000"


def build() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = SPEC["sheetName"]
    columns = SPEC["columns"]
    first_data_row = SPEC["firstDataRow"]

    # 1行目: 見出し
    for i, col in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=i, value=col["header"])
        cell.font = Font(name=FONT, bold=True, size=10)
        cell.fill = REQUIRED_FILL if col["required"] else HEADER_FILL
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = col["width"]

    # 2行目: 記入例
    for i, col in enumerate(columns, start=1):
        cell = ws.cell(row=2, column=i, value=col["example"])
        cell.font = Font(name=FONT, size=10, italic=True, color="84837D")
        cell.fill = EXAMPLE_FILL
        cell.border = BORDER
        if col["kind"] == "yen":
            cell.number_format = YEN_FMT
        elif col["kind"] == "percent":
            cell.number_format = PCT_FMT

    ws.row_dimensions[1].height = 34

    # 3行目以降: 入力欄。書式とプルダウンを先に敷いておく
    last_row = first_data_row + 199
    for i, col in enumerate(columns, start=1):
        letter = get_column_letter(i)
        for row in range(first_data_row, last_row + 1):
            cell = ws.cell(row=row, column=i)
            cell.font = Font(name=FONT, size=10)
            cell.border = BORDER
            if col["kind"] == "yen":
                cell.number_format = YEN_FMT
            elif col["kind"] == "percent":
                cell.number_format = PCT_FMT

        if col["kind"] == "structure":
            options = ",".join(SPEC["structureLabels"])
            dv = DataValidation(
                type="list", formula1=f'"{options}"', allow_blank=True, showDropDown=False
            )
            dv.error = "一覧から選んでください。表記が違うと取り込めません。"
            dv.errorTitle = "構造"
            ws.add_data_validation(dv)
            dv.add(f"{letter}{first_data_row}:{letter}{last_row}")

        if col["kind"] == "taxRate":
            options = ",".join(str(r) for r in SPEC["taxRates"])
            dv = DataValidation(
                type="list", formula1=f'"{options}"', allow_blank=True, showDropDown=False
            )
            dv.error = "33 / 40 / 45 のいずれかを選んでください。"
            dv.errorTitle = "所得税率(%)"
            ws.add_data_validation(dv)
            dv.add(f"{letter}{first_data_row}:{letter}{last_row}")

    ws.freeze_panes = f"A{first_data_row}"

    # 使い方シート
    guide = wb.create_sheet("使い方")
    lines = [
        ("不動産投資提案書 入力テンプレート", True),
        ("", False),
        (f"1. 「{SPEC['sheetName']}」シートの {first_data_row} 行目から、1行に1物件ずつ入力します。", False),
        ("2. 2行目は記入例です。消しても、そのまま残しても構いません(取り込み時は無視されます)。", False),
        ("3. 見出しが薄い黄色の列は必須です。空欄のままだと、その行は取り込まれません。", False),
        ("4. 「構造」と「所得税率(%)」はセルを選ぶとプルダウンが出ます。手入力すると表記違いで弾かれます。", False),
        ("5. 金額は円単位、税抜きや万円ではなくそのままの数字で入れてください(例: 3,000万円 → 30000000)。", False),
        ("6. 「ローン金利(%)」は 2.5% なら 2.5 と入れます。0.025 ではありません。", False),
        ("7. 「うち設備価格」はRC造のときだけ使われます。それ以外の構造では無視されます。", False),
        ("8. 「管理費・修繕積立金(月額)」は管理費と修繕積立金を合算した月額です。", False),
        ("9. 写真はExcelでは扱いません。取り込んだあと、アプリの入力画面で1枚だけ追加してください。", False),
        ("", False),
        ("入力し終えたら、アプリの一覧画面で「Excelから取り込み」を押してこのファイルを選びます。", False),
        ("", False),
        ("列の意味で迷ったら docs/spec.md の第4章を参照してください。", False),
    ]
    for r, (text, is_title) in enumerate(lines, start=1):
        cell = guide.cell(row=r, column=1, value=text)
        cell.font = Font(name=FONT, size=13 if is_title else 10, bold=is_title)
        cell.alignment = Alignment(vertical="center")
    guide.column_dimensions["A"].width = 100

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"wrote {OUT.relative_to(ROOT)}  ({len(columns)} columns, data from row {first_data_row})")


if __name__ == "__main__":
    build()
